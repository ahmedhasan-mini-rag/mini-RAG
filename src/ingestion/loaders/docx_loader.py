"""DOCX loader using *python-docx*.

Extracts text paragraph-by-paragraph.  When a heading style is
detected (``Heading 1``, ``Heading 2``, …), the content under that
heading is grouped into a section and emitted as a separate
:class:`LoadedDocument` with a ``section`` metadata key.  If no
headings are found, the entire document is returned as a single
:class:`LoadedDocument`.
"""

from __future__ import annotations

import asyncio
import logging
from pathlib import Path

from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from docx.opc.exceptions import PackageNotFoundError

from exceptions import FileIOError
from .base import BaseLoader, register_loader
from .schemas import LoadedDocument, DocumentMetaData, ProcessingOutputType

logger = logging.getLogger(__name__)


@register_loader('.docx')
class DocxLoader(BaseLoader):
    """Extract text from ``.docx`` files with optional heading-based
    section splitting."""

    async def load(self, file_path: Path) -> list[LoadedDocument]:
        try:
            doc = await asyncio.to_thread(DocxDocument, str(file_path))
        except PackageNotFoundError as exc:
            raise FileIOError(
                f"Failed to open DOCX file '{file_path.name}'. File may be corrupted or not a valid DOCX",
                detail=str(exc),
            ) from exc
        except Exception as exc:
            raise FileIOError(
                f"Failed to read DOCX file '{file_path.name}'",
                detail=str(exc),
            ) from exc

        sections = self._split_by_headings(doc, file_path)

        if sections:
            return sections

        # if no headings found return entire document as one LoadedDocument
        full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if not full_text.strip():
            return []

        return [
            LoadedDocument(
                text=full_text,
                metadata=DocumentMetaData(
                    doc_type=ProcessingOutputType.TEXT,
                    source=str(file_path),
                    format=".docx"
                )
            )
        ]

    def _split_by_headings(
        self,
        doc: DocxDocumentType,
        file_path: Path,
    ) -> list[LoadedDocument]:
        """Return one LoadedDocument per heading section, or an empty
        list when the document contains no headings."""

        documents: list[LoadedDocument] = []
        current_heading: str | None = None
        current_level: int | None = None
        current_paragraphs: list[str] = []
        has_headings = False

        def _flush():
            """Append the accumulated section to *documents*."""
            if current_heading is None:
                return

            parts = [current_heading] + current_paragraphs
            section_text = "\n".join(parts)

            documents.append(
                LoadedDocument(
                    text=section_text,
                    metadata=DocumentMetaData(
                        doc_type=ProcessingOutputType.TEXT,
                        source=str(file_path),
                        format=".docx",
                        section=current_heading,
                        section_level=current_level
                    )
                )
            )

        for para in doc.paragraphs:
            style_name = (para.style.name or "").lower() # type: ignore

            if style_name.startswith("heading "):
                has_headings = True
                _flush()

                current_heading = para.text.strip()
                try:
                    current_level = int(style_name.split()[-1])
                except (ValueError, IndexError):
                    current_level = None
                current_paragraphs = []
            else:
                if para.text.strip():
                    current_paragraphs.append(para.text)

        # flush the last section
        _flush()

        if not has_headings:
            return []

        return documents

